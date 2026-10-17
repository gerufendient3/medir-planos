"""Punto 4: precios. Plantilla vacia y marcado visible del total incompleto.

Regla dura: el motor NUNCA propone un precio. La plantilla sale con todos los
precios en null y cada concepto sin precio se excluye del subtotal.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.exportar import exportar_excel, exportar_pdf
from app.medidor import medir
from app.presupuesto import Catalogo, construir_presupuesto

RAIZ = Path(__file__).resolve().parent.parent
PLANO = RAIZ / "planos" / "plano_realista.dxf"
PLANTILLA = RAIZ / "catalogo" / "mi_precios.json"


def _texto_pdf(ruta: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)


# ------------------------------------------------------------- plantilla

def test_la_plantilla_se_genera_con_todos_los_precios_en_null(tmp_path):
    import sys

    sys.path.insert(0, str(RAIZ / "scripts"))
    import generar_plantilla_precios as gpp

    gpp.DESTINO = tmp_path / "plantilla.json"
    gpp.main()

    datos = json.loads((tmp_path / "plantilla.json").read_text(encoding="utf-8"))
    assert datos["_meta"]["_plantilla"] is True
    assert all(c["precio_unitario"] is None for c in datos["conceptos"])


def test_la_plantilla_no_contiene_precios_numericos(tmp_path):
    """Regresion contra la filtracion de precios inventados."""
    import sys

    sys.path.insert(0, str(RAIZ / "scripts"))
    import generar_plantilla_precios as gpp

    gpp.DESTINO = tmp_path / "plantilla.json"
    gpp.main()

    datos = json.loads((tmp_path / "plantilla.json").read_text(encoding="utf-8"))
    for c in datos["conceptos"]:
        assert c["precio_unitario"] is None, "{} trae precio".format(c["clave"])


def test_catalogo_de_sin_precio_da_subtotal_cero():
    """Sin precios el subtotal es 0.00 y eso es correcto, no un bug."""
    cat = Catalogo.cargar()
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)
    assert p.subtotal == 0.0
    assert p.iva == 0.0
    assert p.total == 0.0
    assert not p.completo


def test_cada_concepto_sin_precio_queda_marcado():
    cat = Catalogo.cargar()
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)

    sin_precio = p.falta_precio
    assert sin_precio, "deben existir lineas sin precio con el catalogo vacio"
    for l in sin_precio:
        assert l.estado == "FALTA_PRECIO"
        assert l.precio_unitario is None
        assert l.importe is None


def test_el_total_se_marca_incompleto(tmp_path):
    cat = Catalogo.cargar()
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)
    assert p.total_incompleto is True
    assert any("NO ENVIAR" in a for a in p.advertencias())


# ------------------------------------------------- marcado en los entregables

def test_excel_marca_el_total_como_incompleto(tmp_path):
    cat = Catalogo.cargar()
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)
    xlsx = exportar_excel(p, r, tmp_path / "p.xlsx")

    wb = load_workbook(xlsx)
    ws = wb.active
    texto = "\n".join(
        str(ws.cell(row=row, column=col).value or "")
        for row in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
    )
    assert "TOTAL INCOMPLETO" in texto
    assert "NO ES UN PRESUPUESTO ENVIABLE" in texto


def test_pdf_avisa_que_no_se_envie_al_cliente(tmp_path):
    cat = Catalogo.cargar()
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)
    pdf = exportar_pdf(p, r, tmp_path / "p.pdf")
    texto = _texto_pdf(pdf)
    assert "TOTAL INCOMPLETO" in texto
    assert "NO ENVIAR AL CLIENTE" in texto


def test_con_todos_los_precios_no_quedan_lineas_sin_precio(tmp_path):
    """Si el usuario carga TODOS los precios, no queda nada por precio.

    Ojo: el total puede seguir incompleto por otras razones (unidad
    incompatible, capa sin concepto). Eso es correcto y es otro indicador.
    """
    datos = json.loads((RAIZ / "catalogo" / "conceptos.json").read_text(encoding="utf-8"))
    # Precio de ejemplo, no de mercado.
    for c in datos["conceptos"]:
        c["precio_unitario"] = 100.0
    ruta = tmp_path / "todo_precio.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")

    cat = Catalogo.cargar(ruta)
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)

    assert not p.falta_precio, "con todos los precios no debe haber FALTA_PRECIO"
    assert not any("sin precio_unitario" in a for a in p.advertencias())


def test_el_total_incompleto_tiene_causa_distinta_a_los_precios(tmp_path):
    """El plano realista tiene lineas incompatibles aunque todos tengan precio.

    Eso explica por que total_incompleto sigue en True: no es por falta de
    precios, es por geometria que no corresponde al concepto.
    """
    datos = json.loads((RAIZ / "catalogo" / "conceptos.json").read_text(encoding="utf-8"))
    for c in datos["conceptos"]:
        c["precio_unitario"] = 100.0
    ruta = tmp_path / "p.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")

    cat = Catalogo.cargar(ruta)
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)

    assert not p.falta_precio
    assert p.incompatibles, "deben quedar lineas incompatibles en este plano"
    assert p.total_incompleto is True
    assert any("unidad incompatible" in a for a in p.advertencias())


def test_el_total_no_es_cero_si_hay_precios(tmp_path):
    datos = json.loads((RAIZ / "catalogo" / "conceptos.json").read_text(encoding="utf-8"))
    for c in datos["conceptos"]:
        c["precio_unitario"] = 100.0
    ruta = tmp_path / "p.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")

    cat = Catalogo.cargar(ruta)
    r = medir(PLANO)
    p = construir_presupuesto(r["mediciones"], cat)
    assert p.subtotal > 0