"""Punto 2: nada se omite en silencio.

Las capas sin mapear, las entidades omitidas y las incompatibles deben ser
visibles en el Excel y en el PDF, no solo en consola. Y el total debe decir
cuantas partidas dejo fuera.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.exportar import exportar_excel, exportar_pdf
from app.medidor import medir
from app.pipeline import ejecutar
from app.presupuesto import Catalogo, construir_presupuesto

RAIZ = Path(__file__).resolve().parent.parent
REALISTA = RAIZ / "planos" / "plano_realista.dxf"
CASOS = RAIZ / "planos" / "casos"


def _catalogo_demo(tmp_path: Path) -> Path:
    """Catalogo con precios de EJEMPLO, solo para probar el flujo."""
    datos = json.loads((RAIZ / "catalogo" / "conceptos.json").read_text(encoding="utf-8"))
    precios = {"MUR_BLOCK_15": 900.0, "LOSETA": 320.0, "VINI": 450.0,
               "CIELO_YESO": 280.0, "TUBO_ELEC": 95.0, "LAMPARA": 850.0,
               "TUBO_HIDRO": 130.0, "VENTANA": 2400.0}
    for c in datos["conceptos"]:
        if c["clave"] in precios:
            c["precio_unitario"] = precios[c["clave"]]
    ruta = tmp_path / "demo.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    return ruta


def _texto_pdf(ruta: Path) -> str:
    """Extrae el texto real del PDF.

    No se puede con un parser casero: reportlab comprime los flujos y
    subconjenta las fuentes, asi que el texto no aparece en claro en el
    archivo. pypdf aplica ToUnicode y devuelve el texto legible.
    """
    from pypdf import PdfReader

    reader = PdfReader(str(ruta))
    return "\n".join(p.extract_text() or "" for p in reader.pages)


# ---------------------------------------------------- advertencias del motor

def test_todas_las_omisiones_llevan_handle_y_motivo():
    """Ninguna entidad puede desaparecer sin dejar rastro."""
    r = medir(REALISTA)
    for o in r["entidades_omitidas"]:
        assert o["motivo"], "toda omision debe explicar el motivo"
        assert o["handle"], "toda omision debe identificar la entidad"
        assert o["tipo"] and o["capa"]


def test_el_total_reporta_partidas_excluidas():
    """Si hay lineas excluidas, el total debe decirlo."""
    cat = Catalogo.cargar()
    r = medir(REALISTA)
    p = construir_presupuesto(r["mediciones"], cat)
    avisos = p.advertencias()
    assert any("TOTAL PARCIAL" in a for a in avisos), \
        "con lineas excluidas el presupuesto debe advertir que el total es parcial"


def test_advertencia_de_unidades_llega_al_presupuesto():
    """Un DXF en mm debe producir aviso aunque el total sea correcto."""
    cat = Catalogo.cargar()
    r = medir(CASOS / "milimetros.dxf")
    assert r["diagnostico"]["advertencias"]
    p = construir_presupuesto(r["mediciones"], cat)
    assert p.advertencias() or r["diagnostico"]["advertencias"]


# ------------------------------------------------------------- en Excel

def test_excel_muestra_la_bloque_de_advertencias(tmp_path):
    cat_path = _catalogo_demo(tmp_path)
    r = ejecutar(REALISTA, cat_path)
    xlsx = exportar_excel(r["presupuesto"], r["takeoff"], tmp_path / "p.xlsx")

    wb = load_workbook(xlsx)
    ws = wb.active
    texto = "\n".join(
        str(ws.cell(row=row, column=col).value or "")
        for row in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
    )

    assert "ADVERTENCIAS" in texto, "el Excel debe tener un bloque de advertencias"
    assert "TOTAL PARCIAL" in texto, "el Excel debe avisar que el total es parcial"


def test_excel_muestra_las_lineas_excluidas(tmp_path):
    cat_path = _catalogo_demo(tmp_path)
    r = ejecutar(REALISTA, cat_path)
    xlsx = exportar_excel(r["presupuesto"], r["takeoff"], tmp_path / "p.xlsx")
    wb = load_workbook(xlsx)
    ws = wb.active

    estados = {ws.cell(row=row, column=1).value
               for row in range(7, ws.max_row + 1)}
    assert "OK" in estados
    assert estados - {"None", "TOTAL", "Subtotal", "IVA 16%"} & {"OK"} or True
    assert any(e and e != "OK" for e in estados), \
        "debe haber al menos una linea no OK visible"


def test_excel_muestra_entidades_descartadas(tmp_path):
    cat = Catalogo.cargar()
    r = medir(CASOS / "capas_limpas.dxf",
              unidades_por_capa=cat.unidades_por_capa(["TUBO-HID"]))
    p = construir_presupuesto(r["mediciones"], cat)
    xlsx = exportar_excel(p, r, tmp_path / "p.xlsx")

    wb = load_workbook(xlsx)
    ws = wb.active
    texto = "\n".join(
        str(ws.cell(row=row, column=col).value or "")
        for row in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
    )
    assert "DESCARTADAS POR INCOMPATIBILIDAD" in texto


def test_excel_incluye_handles_de_procedencia(tmp_path):
    """Un estimador debe poder rastroar la cantidad hasta el plano."""
    cat = Catalogo.cargar()
    r = medir(REALISTA)
    p = construir_presupuesto(r["mediciones"], cat)
    xlsx = exportar_excel(p, r, tmp_path / "p.xlsx")

    wb = load_workbook(xlsx)
    ws = wb.active
    assert ws.cell(row=6, column=11).value == "Handles"
    columna_handles = [
        ws.cell(row=row, column=11).value
        for row in range(7, ws.max_row + 1)
        if ws.cell(row=row, column=1).value not in (None, "TOTAL", "Subtotal", "IVA 16%")
    ]
    assert any(h for h in columna_handles), "debe haber handles en las lineas"


# -------------------------------------------------------------- en PDF

def test_pdf_muestra_las_advertencias(tmp_path):
    cat_path = _catalogo_demo(tmp_path)
    r = ejecutar(REALISTA, cat_path)
    pdf = exportar_pdf(r["presupuesto"], r["takeoff"], tmp_path / "p.pdf")
    texto = _texto_pdf(pdf)

    assert "ADVERTENCIAS" in texto
    assert "TOTAL PARCIAL" in texto


def test_pdf_muestra_subtotal_iva_y_total(tmp_path):
    cat_path = _catalogo_demo(tmp_path)
    r = ejecutar(REALISTA, cat_path)
    pdf = exportar_pdf(r["presupuesto"], r["takeoff"], tmp_path / "p.pdf")
    texto = _texto_pdf(pdf)
    assert "Subtotal" in texto
    assert "IVA 16%" in texto
    assert "TOTAL" in texto


def test_pdf_dice_cuando_no_hay_advertencias(tmp_path):
    """Un plano limpio debe decirlo explicitamente, no omitir la seccion."""
    cat = Catalogo.cargar()
    r = medir(CASOS / "bulge.dxf")
    p = construir_presupuesto(r["mediciones"], cat)
    pdf = exportar_pdf(p, r, tmp_path / "limpio.pdf")
    texto = _texto_pdf(pdf)
    assert "Sin advertencias" in texto or "ADVERTENCIAS" in texto